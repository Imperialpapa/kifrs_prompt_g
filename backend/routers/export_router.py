"""
Export Router - 파일 다운로드 및 내보내기
"""

import io
import traceback
from datetime import datetime
from urllib.parse import quote

import pandas as pd
from fastapi import APIRouter, File, UploadFile, HTTPException, Form
from fastapi.responses import StreamingResponse

from models import ValidationResponse
from dependencies import fix_service, validation_service
from utils.excel_parser import sanitize_sheet_name
from utils.logger import get_logger

logger = get_logger("export_router")
router = APIRouter(tags=["Export"])


def _safe_str(value, max_length=32000):
    """엑셀 안전 문자열 변환"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    s = str(value)
    if len(s) > max_length:
        return s[:max_length] + "..."
    return s


@router.post("/download-results")
async def download_validation_results(validation_response: ValidationResponse):
    """검증 결과를 Excel 파일로 다운로드"""
    try:
        output = io.BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # 1. 요약 시트
            summary_data = {
                "항목": ["검증 상태", "전체 행 수", "정상 행 수", "오류 행 수", "총 오류 수", "적용된 규칙 수", "검증 시각"],
                "값": [
                    validation_response.validation_status,
                    validation_response.summary.total_rows,
                    validation_response.summary.valid_rows,
                    validation_response.summary.error_rows,
                    validation_response.summary.total_errors,
                    validation_response.summary.rules_applied,
                    validation_response.summary.timestamp.strftime("%Y-%m-%d %H:%M:%S")
                ]
            }
            df_summary = pd.DataFrame(summary_data)
            df_summary.to_excel(writer, sheet_name=sanitize_sheet_name("검증 요약"), index=False)

            # 2. 인지 항목 집계 시트
            if validation_response.error_groups:
                groups_data = []
                for group in validation_response.error_groups:
                    rows_str = ', '.join(map(str, group.affected_rows[:20]))
                    if len(group.affected_rows) > 20:
                        rows_str += f" 외 {len(group.affected_rows) - 20}개"

                    groups_data.append({
                        "시트명": _safe_str(group.sheet),
                        "열": _safe_str(group.column),
                        "규칙ID": _safe_str(group.rule_id),
                        "인지 메시지": _safe_str(group.message, 1000),
                        "인지 횟수": group.count,
                        "영향받은 행": _safe_str(rows_str, 500),
                        "샘플 값": _safe_str(", ".join(map(str, group.sample_values[:10])), 500),
                        "예상 값": _safe_str(group.expected),
                        "원본 규칙": _safe_str(group.source_rule, 500)
                    })
                df_groups = pd.DataFrame(groups_data)
                df_groups.to_excel(writer, sheet_name=sanitize_sheet_name("인지 항목 집계"), index=False)

            # 3. 개별 인지 목록 시트
            if validation_response.errors:
                errors_data = []
                for error in validation_response.errors:
                    errors_data.append({
                        "시트명": _safe_str(error.sheet),
                        "행": error.row,
                        "열": _safe_str(error.column),
                        "규칙ID": _safe_str(error.rule_id),
                        "인지 메시지": _safe_str(error.message, 1000),
                        "실제 값": _safe_str(error.actual_value, 500),
                        "예상 값": _safe_str(error.expected),
                        "원본 규칙": _safe_str(error.source_rule, 500)
                    })
                df_errors = pd.DataFrame(errors_data)
                df_errors.to_excel(writer, sheet_name=sanitize_sheet_name("개별 인지 목록"), index=False)

            # 4. 규칙 충돌 시트
            if validation_response.conflicts:
                conflicts_data = []
                for conflict in validation_response.conflicts:
                    conflicts_data.append({
                        "규칙ID": _safe_str(conflict.rule_id),
                        "충돌 유형": _safe_str(conflict.conflict_type),
                        "설명": _safe_str(conflict.description, 1000),
                        "K-IFRS 1019 참조": _safe_str(conflict.kifrs_reference),
                        "영향받는 규칙": _safe_str(", ".join(conflict.affected_rules), 500),
                        "권장사항": _safe_str(conflict.recommendation, 1000),
                        "심각도": _safe_str(conflict.severity)
                    })
                df_conflicts = pd.DataFrame(conflicts_data)
                df_conflicts.to_excel(writer, sheet_name=sanitize_sheet_name("규칙 충돌"), index=False)

            # 5. 적용된 규칙 시트
            if validation_response.rules_applied:
                rules_data = []
                for rule in validation_response.rules_applied:
                    rules_data.append({
                        "규칙ID": _safe_str(rule.rule_id),
                        "필드명": _safe_str(rule.field_name),
                        "규칙 유형": _safe_str(rule.rule_type),
                        "파라미터": _safe_str(rule.parameters, 500),
                        "오류 메시지": _safe_str(rule.error_message_template, 500),
                        "원본 규칙": _safe_str(rule.source.original_text, 500),
                        "신뢰도": rule.confidence_score
                    })
                df_rules = pd.DataFrame(rules_data)
                df_rules.to_excel(writer, sheet_name=sanitize_sheet_name("적용된 규칙"), index=False)

        timestamp = datetime.now().strftime("%Y-%m-%d")
        filename = f"DBO_Validation_Result_{timestamp}.xlsx"
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-cache"
            }
        )

    except Exception as e:
        logger.error(f"Download results error: {e}")
        raise HTTPException(
            status_code=500,
            detail={"error": "Failed to generate Excel file", "message": str(e)}
        )


@router.post("/api/compare-files")
async def compare_files(
    file1: UploadFile = File(..., description="이전 시점 파일"),
    file2: UploadFile = File(..., description="현재 시점 파일")
):
    """시계열 파일 비교 검증"""
    try:
        logger.info(f"Comparing files: {file1.filename} vs {file2.filename}")
        from services.comparison_service import ComparisonService
        comparison_service = ComparisonService()

        content1 = await file1.read()
        content2 = await file2.read()

        result = comparison_service.compare_files(
            file1_content=content1,
            file2_content=content2,
            file1_name=file1.filename or "파일1",
            file2_name=file2.filename or "파일2"
        )

        logger.info(f"Compare result: {result['summary'].get('description', '')}")
        return result

    except Exception as e:
        logger.error(f"File comparison error: {e}")
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={"error": "File comparison failed", "message": str(e)}
        )


@router.post("/api/export/highlighted")
async def export_highlighted_excel(
    employee_file: UploadFile = File(..., description="직원 데이터 파일"),
    session_id: str = Form(..., description="검증 세션 ID")
):
    """검증 결과 Excel 내보내기 - 원본에 오류 셀 하이라이팅"""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill, Font, Border, Side
        from openpyxl.comments import Comment

        content = await employee_file.read()
        session_details = await validation_service.get_session_details(session_id)

        if not session_details:
            raise HTTPException(status_code=404, detail="Session not found")

        errors = session_details.get("errors", [])

        wb = load_workbook(io.BytesIO(content))

        error_map = {}
        for err in errors:
            sheet = err.get("sheet", "")
            row = err.get("row", 0)
            col = err.get("column", "")
            msg = err.get("message", "")
            key = (sheet, row, col)
            if key not in error_map:
                error_map[key] = []
            error_map[key].append(msg)

        error_fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
        error_font = Font(color="CC0000")
        error_border = Border(
            left=Side(style='thin', color='CC0000'),
            right=Side(style='thin', color='CC0000'),
            top=Side(style='thin', color='CC0000'),
            bottom=Side(style='thin', color='CC0000')
        )

        for ws in wb.worksheets:
            sheet_name = ws.title
            header_row = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}

            for (err_sheet, err_row, err_col), messages in error_map.items():
                if err_sheet != sheet_name:
                    continue
                if err_col in header_row and err_row > 0:
                    col_idx = header_row[err_col]
                    cell = ws.cell(row=err_row, column=col_idx)
                    cell.fill = error_fill
                    cell.font = error_font
                    cell.border = error_border
                    comment_text = "\n".join(messages[:3])
                    cell.comment = Comment(comment_text, "DBO Validator")

        if "오류요약" in wb.sheetnames:
            del wb["오류요약"]
        summary_ws = wb.create_sheet("오류요약", 0)
        summary_ws.append(["시트", "행", "컬럼", "오류 메시지", "실제 값", "예상 값"])
        for err in errors[:500]:
            summary_ws.append([
                err.get("sheet", ""),
                err.get("row", 0),
                err.get("column", ""),
                err.get("message", ""),
                str(err.get("actual_value", "")),
                str(err.get("expected", ""))
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        base_name = (employee_file.filename or "data").rsplit('.', 1)[0]
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{base_name}_highlighted_{timestamp}.xlsx"
        filename_encoded = quote(filename, safe='')

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{filename_encoded}",
                "Cache-Control": "no-cache"
            }
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Highlighted export error: {e}")
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={"error": "Highlighted export failed", "message": str(e)}
        )


@router.post("/api/fix/download")
async def bulk_fix_download(
    cells_to_fix_json: str = Form(..., description="JSON array of cells to fix"),
    original_file: UploadFile = File(..., description="Original Excel file")
):
    """오류 항목 일괄 수정 후 엑셀 파일 다운로드"""
    try:
        import json
        cells_to_fix = json.loads(cells_to_fix_json)

        logger.info(f"Bulk fix download: {len(cells_to_fix)} cells from {original_file.filename}")

        content = await original_file.read()

        modified_excel = fix_service.apply_bulk_fixes_to_excel(
            content,
            cells_to_fix,
            filename=original_file.filename or ""
        )

        base_name = original_file.filename.rsplit('.', 1)[0] if original_file.filename else "data"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{base_name}_fixed_{timestamp}.xlsx"
        filename_encoded = quote(filename, safe='')

        return StreamingResponse(
            io.BytesIO(modified_excel),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{filename_encoded}",
                "Content-Length": str(len(modified_excel)),
                "Cache-Control": "no-cache"
            }
        )

    except Exception as e:
        logger.error(f"Bulk fix download error: {e}")
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={"error": "Failed to create fixed file", "message": str(e)}
        )
